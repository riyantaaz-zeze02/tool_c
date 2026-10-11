import warnings
import subprocess
import sys

import openpyxl
import pandas as pd

from cleaner.engine import CleaningEngine


def _write_csv(path, df):
    df.to_csv(path, index=False)


def test_append_keeps_base_untouched_and_adds_only_new_rows(tmp_path):
    base = pd.DataFrame({
        "Order ID": ["A1", "A2"],
        "Nama": ["Budi", "Siti"],
        "Harga": ["Rp 100.000", "Rp 200.000"],
        "Tanggal": ["15/01/2025", "20/01/2025"],
    })
    new = pd.DataFrame({
        "Order ID": ["A2", "A3"],
        "Nama": ["Siti", "Andi"],
        "Harga": ["Rp 200.000", "Rp 300.000"],
        "Tanggal": ["20/01/2025", "05/02/2025"],
    })
    base_path = tmp_path / "base.csv"
    new_path = tmp_path / "new.csv"
    _write_csv(base_path, base)
    _write_csv(new_path, new)

    df_final, report = CleaningEngine().append_files(str(base_path), [str(new_path)])

    assert len(df_final) == 3
    assert report["total_ditambahkan"] == 1
    assert report["total_dilewati"] == 1
    assert list(df_final["Order ID"]) == ["A1", "A2", "A3"]


def test_append_reorders_columns_and_skips_total_row(tmp_path):
    base = pd.DataFrame({
        "Order ID": ["A1", "A2"],
        "Nama": ["Budi", "Siti"],
        "Harga": ["Rp 100.000", "Rp 200.000"],
        "Tanggal": ["15/01/2025", "20/01/2025"],
    })
    base.loc[len(base)] = ["TOTAL", "", "", ""]
    new = pd.DataFrame({
        "Nama": ["Andi"],
        "Tanggal": ["05/02/2025"],
        "Harga": ["Rp 300.000"],
        "Order ID": ["A3"],
    })
    base_path = tmp_path / "base.csv"
    new_path = tmp_path / "new.csv"
    _write_csv(base_path, base)
    _write_csv(new_path, new)

    df_final, report = CleaningEngine().append_files(str(base_path), [str(new_path)])

    assert df_final["Order ID"].tolist() == ["A1", "A2", "A3"]
    assert report["base_total_row_dibuang"] == 1
    assert report["total_ditambahkan"] == 1
    assert report["baris_akhir"] == 3


def test_append_rejects_different_column_set(tmp_path):
    base = pd.DataFrame({"A": [1], "B": [2]})
    new = pd.DataFrame({"A": [3], "C": [4]})
    base_path = tmp_path / "base.csv"
    new_path = tmp_path / "new.csv"
    _write_csv(base_path, base)
    _write_csv(new_path, new)

    try:
        CleaningEngine().append_files(str(base_path), [str(new_path)])
        assert False, "Expected ValueError"
    except ValueError as exc:
        assert "Kolom hilang" in str(exc) or "kolom" in str(exc).lower()


def test_append_uses_dedupe_keys_to_skip_rows(tmp_path):
    base = pd.DataFrame({
        "Order ID": ["A1"],
        "Nama": ["Budi"],
        "Harga": ["Rp 100.000"],
    })
    new = pd.DataFrame({
        "Order ID": ["A1", "A3", "A3", "A2"],
        "Nama": ["Budi", "Cici", "Cici", "Ayu"],
        "Harga": ["Rp 100.000", "Rp 999.000", "Rp 900.000", "Rp 150.000"],
    })
    base_path = tmp_path / "base.csv"
    new_path = tmp_path / "new.csv"
    _write_csv(base_path, base)
    _write_csv(new_path, new)

    df_final, report = CleaningEngine().append_files(str(base_path), [str(new_path)], dedupe_keys=["Order ID"])

    assert list(df_final["Order ID"]) == ["A1", "A3", "A2"]
    assert report["total_dilewati"] == 2
    skipped = report["skipped_rows"]
    assert set(skipped["Alasan"].tolist()) == {"sudah ada di file dasar", "duplikat di file baru"}


def test_append_similar_ids_are_not_coerced_to_numbers(tmp_path):
    base_path = tmp_path / "base.csv"
    new_path = tmp_path / "new.csv"
    _write_csv(base_path, pd.DataFrame({"Order ID": ["ORD-10641"]}))
    _write_csv(new_path, pd.DataFrame({"Order ID": ["INV-10641", "ABC-10641"]}))

    final, report = CleaningEngine().append_files(
        str(base_path), [str(new_path)], dedupe_keys=["Order ID"]
    )

    assert final["Order ID"].tolist() == ["ORD-10641", "INV-10641", "ABC-10641"]
    assert report["total_dilewati"] == 0


def test_compare_coercion_requires_a_full_numeric_value():
    engine = CleaningEngine()

    assert engine._coerce_for_compare("ORD-10641") == "ord-10641"
    assert engine._coerce_for_compare("100€") == 100.0
    assert engine._coerce_for_compare("Rp 1.234.567,00*") == 1234567.0


def test_append_identical_ambiguous_dates_are_deduplicated_without_warning(tmp_path):
    base_path = tmp_path / "base.csv"
    new_path = tmp_path / "new.csv"
    _write_csv(base_path, pd.DataFrame({"Order ID": ["A1"], "Tanggal": ["03/04/2025"]}))
    _write_csv(new_path, pd.DataFrame({"Order ID": ["A1"], "Tanggal": ["03/04/2025"]}))

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        final, report = CleaningEngine().append_files(str(base_path), [str(new_path)])

    assert len(final) == 1
    assert report["total_dilewati"] == 1
    assert not [warning for warning in caught if "dayfirst" in str(warning.message)]


def test_append_windows_1252_csv_and_selected_excel_sheet(tmp_path):
    base_path = tmp_path / "base.csv"
    csv_path = tmp_path / "Jos\u00e9 Pe\u00f1a.csv"
    excel_path = tmp_path / "new.xlsx"
    _write_csv(base_path, pd.DataFrame({"Order ID": ["A1"], "Nama": ["Base"]}))
    pd.DataFrame({"Order ID": ["A2"], "Nama": ["Jos\u00e9 Pe\u00f1a"]}).to_csv(
        csv_path, index=False, encoding="windows-1252"
    )
    with pd.ExcelWriter(excel_path) as writer:
        pd.DataFrame({"Order ID": ["WRONG"], "Nama": ["First"]}).to_excel(
            writer, sheet_name="First", index=False
        )
        pd.DataFrame({"Order ID": ["A3"], "Nama": ["Selected"]}).to_excel(
            writer, sheet_name="Target", index=False
        )

    final, report = CleaningEngine().append_files(
        str(base_path), [str(csv_path), str(excel_path)], sheet_name="Target"
    )

    assert final["Order ID"].tolist() == ["A1", "A2", "A3"]
    assert final["Nama"].tolist()[1:] == ["Jos\u00e9 Pe\u00f1a", "Selected"]
    assert report["baris_awal"] == 3


def test_append_export_includes_append_summary(tmp_path):
    base_path = tmp_path / "base.csv"
    new_path = tmp_path / "new.csv"
    output_path = tmp_path / "appended.xlsx"
    _write_csv(base_path, pd.DataFrame({"Order ID": ["A1"]}))
    _write_csv(new_path, pd.DataFrame({"Order ID": ["A2"]}))
    engine = CleaningEngine()
    final, report = engine.append_files(str(base_path), [str(new_path)])

    engine.export_excel(final, str(output_path), report=report, include_summary_sheet=True)

    summary = openpyxl.load_workbook(output_path)["Cleaning Summary"]
    values = [cell.value for row in summary.iter_rows() for cell in row]
    assert "APPEND SUMMARY" in values
    assert "APPEND FILE DETAILS" in values
    assert "NEW ROWS CLEANING SUMMARY" in values


def test_cli_append_exports_summary_and_tags_sources(tmp_path):
    project_root = __import__("pathlib").Path(__file__).resolve().parents[1]
    base_path = tmp_path / "base.csv"
    new_path = tmp_path / "new.csv"
    output_path = tmp_path / "appended.xlsx"
    _write_csv(base_path, pd.DataFrame({"Order ID": ["A1"]}))
    _write_csv(new_path, pd.DataFrame({"Order ID": ["A2"]}))

    result = subprocess.run(
        [
            sys.executable, str(project_root / "cli.py"), str(base_path), "--append",
            "--new", str(new_path), "--dedupe-keys", "Order ID", "--tag-source",
            "--output", str(output_path),
        ],
        cwd=project_root,
        capture_output=True,
    )

    assert result.returncode == 0, result.stderr.decode("utf-8", errors="replace") + result.stdout.decode("utf-8", errors="replace")
    workbook = openpyxl.load_workbook(output_path, data_only=True)
    data = workbook["Data Bersih"]
    assert data.max_row == 4
    assert data.cell(row=2, column=2).value == "base.csv"
    assert data.cell(row=3, column=2).value == "new.csv"
    summary = workbook["Cleaning Summary"]
    assert summary["B5"].value == "APPEND SUMMARY"


def test_cli_append_new_folder_and_base_overwrite_backup(tmp_path):
    project_root = __import__("pathlib").Path(__file__).resolve().parents[1]
    base_path = tmp_path / "base.csv"
    new_folder = tmp_path / "new_files"
    new_folder.mkdir()
    new_path = new_folder / "new.csv"
    _write_csv(base_path, pd.DataFrame({"Order ID": ["A1"]}))
    _write_csv(new_path, pd.DataFrame({"Order ID": ["A2"]}))
    original_base = pd.read_csv(base_path)

    command = [
        sys.executable, str(project_root / "cli.py"), str(base_path), "--append",
        "--new-folder", str(new_folder), "--format", "csv", "--output", str(base_path),
    ]
    rejected = subprocess.run(command, cwd=project_root, capture_output=True)
    assert rejected.returncode != 0
    assert "--overwrite" in rejected.stdout.decode("utf-8", errors="replace")
    pd.testing.assert_frame_equal(pd.read_csv(base_path), original_base)

    allowed = subprocess.run(
        command + ["--overwrite"], cwd=project_root, capture_output=True
    )
    assert allowed.returncode == 0, allowed.stderr.decode("utf-8", errors="replace") + allowed.stdout.decode("utf-8", errors="replace")
    pd.testing.assert_frame_equal(pd.read_csv(str(base_path) + ".bak"), original_base)
    assert pd.read_csv(base_path)["Order ID"].tolist() == ["A1", "A2"]


def test_cli_append_5000_base_500_new_rows_preserves_prices_and_dates(tmp_path):
    project_root = __import__("pathlib").Path(__file__).resolve().parents[1]
    base_path = tmp_path / "base.csv"
    new_path = tmp_path / "new.csv"
    output_path = tmp_path / "appended.xlsx"
    base_dates = pd.date_range("2024-01-01", periods=5000)
    base = pd.DataFrame({
        "Order ID": [f"ORD-{index:05d}" for index in range(1, 5001)],
        "Harga": [1000 + index for index in range(5000)],
        "Tanggal": base_dates.strftime("%Y-%m-%d"),
    })
    new_dates = base_dates[:100].append(
        pd.date_range(base_dates[-1] + pd.Timedelta(days=1), periods=400)
    )
    new = pd.DataFrame({
        "Order ID": [f"ORD-{index:05d}" for index in range(1, 501)],
        "Harga": [f"Rp {1000 + index:,}" for index in range(500)],
        "Tanggal": new_dates.strftime("%d/%m/%Y"),
    })
    new.loc[100:, "Order ID"] = [f"ORD-{index:05d}" for index in range(5001, 5401)]
    _write_csv(base_path, base)
    _write_csv(new_path, new)

    result = subprocess.run(
        [
            sys.executable, str(project_root / "cli.py"), str(base_path), "--append",
            "--new", str(new_path), "--dedupe-keys", "Order ID", "--no-total",
            "--output", str(output_path),
        ],
        cwd=project_root,
        capture_output=True,
    )

    output_text = result.stdout.decode("utf-8", errors="replace")
    assert result.returncode == 0, result.stderr.decode("utf-8", errors="replace") + output_text
    assert "Ditambahkan: 400" in output_text
    assert "Dilewati: 100" in output_text
    workbook = pd.ExcelFile(output_path)
    assert "Cleaning Summary" in workbook.sheet_names
    final = pd.read_excel(output_path, sheet_name="Data Bersih")
    assert len(final) == 5400
    indexed = final.set_index("Order ID")
    assert indexed.index.is_unique
    for index, (order_id, date_value) in enumerate(zip(base["Order ID"], base["Tanggal"])):
        assert indexed.at[order_id, "Harga"] == 1000 + index
        assert indexed.at[order_id, "Tanggal"] == pd.Timestamp(date_value)
    for index, row in new.iloc[100:].reset_index(drop=True).iterrows():
        assert indexed.at[row["Order ID"], "Harga"] == 1100 + index
        assert indexed.at[row["Order ID"], "Tanggal"] == pd.to_datetime(
            row["Tanggal"], format="%d/%m/%Y"
        )
