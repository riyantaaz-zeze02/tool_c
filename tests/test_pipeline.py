"""
test_pipeline.py — Integration Tests untuk Pipeline Cleaning & Export Excel
"""

import os
import openpyxl
import pandas as pd
import pytest
from cleaner.engine import CleaningEngine
from cleaner.ui_helpers import default_sheet_selection
from cli import analyze_auto_files


def test_full_pipeline_on_sample_csv(tmp_path):
    sample_csv = os.path.join("data", "input", "contoh_data.csv")
    assert os.path.exists(sample_csv)

    engine = CleaningEngine()
    df_clean, report = engine.clean(sample_csv)

    # Validasi hasil cleaning
    assert len(df_clean) == 11
    assert report["baris_awal"] == 15
    assert report["total_duplikat_dihapus"] == 4
    assert df_clean["nama"].iloc[0] == "Hani Wijaya"

    # Validasi ekspor Excel
    out_file = str(tmp_path / "output_test.xlsx")
    engine.export_excel(df_clean, out_file, report=report, theme="emerald_green")
    assert os.path.exists(out_file)

    # Validasi isi file Excel
    wb = openpyxl.load_workbook(out_file)
    assert "Data Bersih" in wb.sheetnames
    assert "Cleaning Summary" in wb.sheetnames

    ws_data = wb["Data Bersih"]
    assert ws_data["A1"].value == "nama"
    assert ws_data["E13"].value == "=SUM(E2:E12)"

    ws_summary = wb["Cleaning Summary"]
    assert ws_summary["B2"].value == "DATA CLEANING AUDIT REPORT"
    assert "YANTTT" in str(ws_summary["B3"].value)
    assert ws_summary["B5"].value == "KEY METRICS"
    assert ws_summary["B13"].value == "COLUMN STATUS DETAILS"


def test_merge_three_files_detects_cross_file_duplicates_and_sources(tmp_path):
    columns = ["id", "nama", "nilai"]
    rows = [
        [[1, "andi", 100], [2, "budi", 200]],
        [[2, "budi", 200], [3, "cici", 300]],
        [[4, "dedi", 400]],
    ]
    files = []
    for month, month_rows in zip(("jan", "feb", "mar"), rows):
        path = tmp_path / f"{month}.csv"
        pd.DataFrame(month_rows, columns=columns).to_csv(path, index=False)
        files.append(str(path))

    engine = CleaningEngine()
    cleaned, report = engine.clean_merged_files(files)

    assert list(cleaned.columns) == ["id", "nama", "nilai", "Sumber File"]
    assert len(cleaned) == 4
    assert report["baris_awal"] == 5
    assert report["duplikat_antar_file"] == 1
    assert report["total_duplikat_dihapus"] == 1
    assert set(cleaned["Sumber File"]) == {"jan.csv", "feb.csv", "mar.csv"}
    assert [item["baris_awal"] for item in report["file_details"]] == [2, 2, 1]


def test_merge_rejects_different_column_structure(tmp_path):
    first = tmp_path / "jan.csv"
    second = tmp_path / "feb.csv"
    pd.DataFrame([[1, "andi"]], columns=["id", "nama"]).to_csv(first, index=False)
    pd.DataFrame([[2, "budi"]], columns=["id", "nama_lengkap"]).to_csv(second, index=False)

    with pytest.raises(ValueError, match=r"Struktur kolom tidak sama.*feb\.csv"):
        CleaningEngine().clean_merged_files([str(first), str(second)])


def test_merge_summary_contains_per_file_breakdown(tmp_path):
    files = []
    for name in ("jan", "feb", "mar"):
        path = tmp_path / f"{name}.csv"
        pd.DataFrame([[1, "produk"]], columns=["id", "nama"]).to_csv(path, index=False)
        files.append(str(path))

    engine = CleaningEngine()
    cleaned, report = engine.clean_merged_files(files)
    output = tmp_path / "merged.xlsx"
    engine.export_excel(cleaned, str(output), report=report)

    summary = openpyxl.load_workbook(output)["Cleaning Summary"]
    values = [cell.value for row in summary.iter_rows() for cell in row]
    assert "Initial Rows" in values
    assert "jan.csv" in values
    assert "  • Duplikat Antar-File (Cross-File)" in values


def test_join_three_files_is_chained_and_reports_match_rate(tmp_path):
    customers = tmp_path / "pelanggan.csv"
    transactions = tmp_path / "transaksi.csv"
    products = tmp_path / "produk.csv"
    pd.DataFrame([[1, "Andi"], [2, "Budi"], [3, "Cici"]], columns=["ID", "Nama"]).to_csv(customers, index=False)
    pd.DataFrame([[1, "P01", 100], [2, "P99", 200], [4, "P01", 400]], columns=["ID", "ID_Produk", "Total"]).to_csv(transactions, index=False)
    pd.DataFrame([["P01", "Laptop"], ["P03", "Mouse"]], columns=["ID_Produk", "Produk"]).to_csv(products, index=False)

    cleaned, report = CleaningEngine().clean_joined_files(
        [str(customers), str(transactions), str(products)],
        ["ID", "ID_Produk"],
    )

    assert len(cleaned) == 3
    assert cleaned["Produk"].iloc[0] == "Laptop"
    assert cleaned["Produk"].iloc[1] == "Unknown"
    assert cleaned["Produk"].iloc[2] == "Unknown"
    assert len(report["join_stages"]) == 2
    assert report["join_stages"][0]["match_rate"] == 66.67
    assert report["join_stages"][1]["matched_rows"] == 1
    assert report["join_stages"][1]["unmatched_rows"] == 2
    assert report["join_stages"][1]["match_rate"] == 33.33
    assert report["total_orphan_rows"] == 2
    assert set(report["orphan_rows"]["Tahap"]) == {1, 2}
    assert "ID 4 tidak ditemukan di pelanggan.csv" in set(report["orphan_rows"]["Alasan"])
    assert "ID_Produk P03 tidak ditemukan di hasil gabungan" in set(report["orphan_rows"]["Alasan"])

    output = tmp_path / "joined.xlsx"
    CleaningEngine().export_excel(cleaned, str(output), report=report)
    summary_values = [
        cell.value
        for row in openpyxl.load_workbook(output)["Cleaning Summary"].iter_rows()
        for cell in row
    ]
    assert "JOIN STAGES" in summary_values
    assert "33.33%" in summary_values
    assert "2 baris tidak ikut ke hasil akhir — lihat sheet 'Baris Ter-drop' untuk detail" in summary_values
    assert "Baris Ter-drop" in openpyxl.load_workbook(output).sheetnames
    orphan_sheet = openpyxl.load_workbook(output)["Baris Ter-drop"]
    orphan_values = [cell.value for row in orphan_sheet.iter_rows() for cell in row]
    assert "ID 4 tidak ditemukan di pelanggan.csv" in orphan_values
    assert "ID_Produk P03 tidak ditemukan di hasil gabungan" in orphan_values


def test_join_single_stage_orphan_is_exported(tmp_path):
    left = tmp_path / "pelanggan.csv"
    right = tmp_path / "transaksi.csv"
    pd.DataFrame([[1, "Andi"]], columns=["ID", "Nama"]).to_csv(left, index=False)
    pd.DataFrame([[1, 100], [4, 400]], columns=["ID", "Total"]).to_csv(right, index=False)

    engine = CleaningEngine()
    cleaned, report = engine.clean_joined_files([str(left), str(right)], ["ID"])
    output = tmp_path / "single_stage_join.xlsx"
    engine.export_excel(cleaned, str(output), report=report)

    orphan_sheet = openpyxl.load_workbook(output)["Baris Ter-drop"]
    assert orphan_sheet["A1"].value == "ID"
    assert orphan_sheet["C1"].value == "Tahap"
    assert orphan_sheet["D1"].value == "File Asal"
    assert orphan_sheet["E1"].value == "Alasan"
    assert orphan_sheet["A2"].value == 4
    assert orphan_sheet["C2"].value == 1
    assert orphan_sheet["E2"].value == "ID 4 tidak ditemukan di pelanggan.csv"


def test_join_rejects_wrong_key_count_before_join(tmp_path):
    files = []
    for index in range(3):
        path = tmp_path / f"file{index}.csv"
        pd.DataFrame([[index]], columns=["ID"]).to_csv(path, index=False)
        files.append(str(path))

    with pytest.raises(ValueError, match=r"3 file membutuhkan 2 key"):
        CleaningEngine().clean_joined_files(files, ["ID"])


def test_join_rejects_key_missing_in_specific_file(tmp_path):
    first = tmp_path / "pelanggan.csv"
    second = tmp_path / "transaksi.csv"
    third = tmp_path / "produk.csv"
    pd.DataFrame([[1]], columns=["ID"]).to_csv(first, index=False)
    pd.DataFrame([[1, "P01"]], columns=["ID", "ID_Produk"]).to_csv(second, index=False)
    pd.DataFrame([["P01"]], columns=["KodeProduk"]).to_csv(third, index=False)

    with pytest.raises(ValueError, match=r"Key 'ID_Produk'.*produk\.csv"):
        CleaningEngine().clean_joined_files([str(first), str(second), str(third)], ["ID", "ID_Produk"])

def test_auto_recommends_merge_for_identical_headers(tmp_path):
    files = []
    for name in ("jan", "feb", "mar"):
        path = tmp_path / f"{name}.csv"
        pd.DataFrame([[1, "produk"]], columns=["ID", "Nama"]).to_csv(path, index=False)
        files.append(str(path))

    analysis = analyze_auto_files(files)

    assert analysis["mode"] == "merge"
    assert analysis["keys"] == []


def test_auto_recommends_join_for_clear_chain(tmp_path):
    first = tmp_path / "pelanggan.csv"
    second = tmp_path / "transaksi.csv"
    third = tmp_path / "produk.csv"
    pd.DataFrame([[1]], columns=["ID"]).to_csv(first, index=False)
    pd.DataFrame([[1, "P01"]], columns=["ID", "ID_Produk"]).to_csv(second, index=False)
    pd.DataFrame([["P01"]], columns=["ID_Produk"]).to_csv(third, index=False)

    analysis = analyze_auto_files([str(first), str(second), str(third)])

    assert analysis["mode"] == "join"
    assert [os.path.basename(path) for path in analysis["chain"]] == [
        "pelanggan.csv", "transaksi.csv", "produk.csv"
    ]
    assert analysis["keys"] == ["ID", "ID_Produk"]


def test_auto_rejects_ambiguous_structure(tmp_path):
    first = tmp_path / "first.csv"
    second = tmp_path / "second.csv"
    third = tmp_path / "third.csv"
    pd.DataFrame([[1, "a"]], columns=["ID", "NilaiA"]).to_csv(first, index=False)
    pd.DataFrame([[1, "b"]], columns=["ID", "NilaiB"]).to_csv(second, index=False)
    pd.DataFrame([[1, "c"]], columns=["ID", "NilaiC"]).to_csv(third, index=False)

    analysis = analyze_auto_files([str(first), str(second), str(third)])

    assert analysis["mode"] == "unknown"
    assert "tidak yakin" in analysis["reason"]


def _create_dummy_3sheet_excel(file_path):
    """Helper untuk membuat file Excel 3 sheet untuk testing."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Penjualan"
    ws1.append(["tanggal", "produk", "harga", "kuantitas"])
    ws1.append(["2024-01-15", "  LAPTOP ASUS [VIP]  ", "Rp 15.000.000", 2])
    ws1.append(["2024-01-15", "  LAPTOP ASUS [VIP]  ", "Rp 15.000.000", 2])
    ws1.append(["2024-02-10", "mouse wireless", "150.000", 5])

    ws2 = wb.create_sheet("Pelanggan")
    ws2.append(["id", "nama_pelanggan", "kota", "status"])
    ws2.append([1, "budi santoso [VIP]", "jakarta", "aktif"])
    ws2.append([1, "budi santoso [VIP]", "jakarta", "aktif"])
    ws2.append([2, "siti aminah", "BANDUNG", "aktif"])
    ws2.append([3, "ahmad dahlan", "surabaya", None])

    ws3 = wb.create_sheet("Ringkasan")
    ws3["A1"] = "LAPORAN EKSEKUTIF"
    ws3["A2"] = "Total Target:"
    ws3["B2"] = "=SUM(1000, 2500, 1500)"
    ws3["A3"] = "Catatan:"
    ws3["B3"] = "Formula dan formatting asli harus utuh"

    wb.save(file_path)
    return file_path


def test_multisheet_process_single_sheet_preserves_others(tmp_path):
    """Skenario 1: Memproses hanya 1 sheet spesifik, sheet lain tetap ada dan utuh."""
    input_file = str(tmp_path / "input_3sheet.xlsx")
    _create_dummy_3sheet_excel(input_file)

    engine = CleaningEngine()
    # Bersihkan hanya sheet 'Penjualan'
    cleaned_results = engine.clean_sheets(input_file, sheets=["Penjualan"])
    assert "Penjualan" in cleaned_results
    assert len(cleaned_results["Penjualan"]["df"]) == 2  # duplikat dihapus dari 3 jadi 2

    out_file = str(tmp_path / "output_single_sheet.xlsx")
    engine.export_excel(
        cleaned_results,
        out_file,
        original_file_path=input_file,
    )

    wb_out = openpyxl.load_workbook(out_file)
    assert "Penjualan" in wb_out.sheetnames
    assert "Pelanggan" in wb_out.sheetnames
    assert "Ringkasan" in wb_out.sheetnames
    assert "Cleaning Summary" in wb_out.sheetnames

    # Pastikan sheet yang tidak disentuh (Ringkasan) formulanya tetap utuh
    assert wb_out["Ringkasan"]["B2"].value == "=SUM(1000, 2500, 1500)"


def test_multisheet_default_selection_is_first_sheet_only():
    assert default_sheet_selection(["Penjualan", "Pelanggan", "Ringkasan"]) == ["Penjualan"]


def test_multisheet_default_workflow_only_cleans_first_sheet(tmp_path):
    input_file = str(tmp_path / "input_default_3sheet.xlsx")
    _create_dummy_3sheet_excel(input_file)

    engine = CleaningEngine()
    cleaned_results = engine.clean_sheets(
        input_file,
        sheets=default_sheet_selection(["Penjualan", "Pelanggan", "Ringkasan"]),
    )
    output_file = str(tmp_path / "output_default_3sheet.xlsx")
    engine.export_excel(cleaned_results, output_file, original_file_path=input_file)

    input_workbook = openpyxl.load_workbook(input_file, data_only=False)
    output_workbook = openpyxl.load_workbook(output_file, data_only=False)
    assert len(cleaned_results) == 1
    assert len(output_workbook["Penjualan"]["A"]) >= len(input_workbook["Penjualan"]["A"])
    assert output_workbook["Pelanggan"]["A2"].value == input_workbook["Pelanggan"]["A2"].value
    assert output_workbook["Ringkasan"]["B2"].value == "=SUM(1000, 2500, 1500)"
    assert output_workbook["Ringkasan"]["B2"].value == input_workbook["Ringkasan"]["B2"].value


def test_multisheet_process_multiple_sheets_and_report(tmp_path):
    """Skenario 2: Memproses beberapa sheet sekaligus dengan audit report per sheet."""
    input_file = str(tmp_path / "input_3sheet.xlsx")
    _create_dummy_3sheet_excel(input_file)

    engine = CleaningEngine()
    # Bersihkan sheet 'Penjualan' dan 'Pelanggan'
    cleaned_results = engine.clean_sheets(input_file, sheets=["Penjualan", "Pelanggan"])
    assert "Penjualan" in cleaned_results
    assert "Pelanggan" in cleaned_results

    out_file = str(tmp_path / "output_multi.xlsx")
    engine.export_excel(
        cleaned_results,
        out_file,
        original_file_path=input_file,
    )

    wb_out = openpyxl.load_workbook(out_file)
    assert "Penjualan" in wb_out.sheetnames
    assert "Pelanggan" in wb_out.sheetnames
    assert "Ringkasan" in wb_out.sheetnames
    assert "Cleaning Summary" in wb_out.sheetnames

    ws_sum = wb_out["Cleaning Summary"]
    assert ws_sum["B2"].value == "DATA CLEANING AUDIT REPORT"
    assert ws_sum["B5"].value == "KEY METRICS"
    # Cek bahwa laporan mencatat baris per-sheet
    sheet_names_in_summary = [ws_sum.cell(row=r, column=3).value for r in [7, 8]]
    assert "Penjualan" in sheet_names_in_summary
    assert "Pelanggan" in sheet_names_in_summary


def test_multisheet_unselected_sheet_untouched(tmp_path):
    """Skenario 3: Memastikan sheet yang tidak dipilih benar-benar tidak berubah di output."""
    input_file = str(tmp_path / "input_3sheet.xlsx")
    _create_dummy_3sheet_excel(input_file)

    engine = CleaningEngine()
    # Proses hanya sheet 'Pelanggan'
    cleaned_results = engine.clean_sheets(input_file, sheets=["Pelanggan"])

    out_file = str(tmp_path / "output_unselected_test.xlsx")
    engine.export_excel(
        cleaned_results,
        out_file,
        original_file_path=input_file,
    )

    wb_in = openpyxl.load_workbook(input_file)
    wb_out = openpyxl.load_workbook(out_file)

    # Validasi seluruh sel di sheet Ringkasan sama persis antara input dan output
    ws_in_ringkasan = wb_in["Ringkasan"]
    ws_out_ringkasan = wb_out["Ringkasan"]

    assert ws_out_ringkasan["A1"].value == ws_in_ringkasan["A1"].value == "LAPORAN EKSEKUTIF"
    assert ws_out_ringkasan["A2"].value == ws_in_ringkasan["A2"].value == "Total Target:"
    assert ws_out_ringkasan["B2"].value == ws_in_ringkasan["B2"].value == "=SUM(1000, 2500, 1500)"
    assert ws_out_ringkasan["A3"].value == ws_in_ringkasan["A3"].value == "Catatan:"
    assert ws_out_ringkasan["B3"].value == ws_in_ringkasan["B3"].value == "Formula dan formatting asli harus utuh"

