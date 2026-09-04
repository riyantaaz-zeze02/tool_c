"""
test_formatter.py — Unit Tests untuk Modul Excel Formatting (openpyxl)
"""

import pandas as pd
import openpyxl
from cleaner.formatter.excel_styler import ExcelStyler
from cleaner.formatter.themes import get_theme, THEMES
from cleaner.formatter.number_formats import get_column_format_mask


def test_theme_retrieval():
    theme = get_theme("corporate_blue")
    assert theme["name"] == "Corporate Blue"
    assert theme["header_fill"] == "1F4E79"

    # Fallback jika nama tidak dikenal
    unknown = get_theme("non_existent_theme")
    assert unknown["name"] == "Corporate Blue"


def test_column_format_mask():
    mask_gaji = get_column_format_mask("gaji_karyawan", "float64")
    assert "Rp" in mask_gaji

    mask_date = get_column_format_mask("tanggal_masuk", "datetime64[ns]")
    assert "yy" in mask_date


def test_excel_styler():
    df = pd.DataFrame({
        "nama": ["Budi", "Ani"],
        "gaji": [5000000.0, 6000000.0],
    })

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Sheet"

    for c_i, col in enumerate(df.columns, 1):
        ws.cell(row=1, column=c_i, value=col)
    for r_i, row in enumerate(df.itertuples(index=False), 2):
        for c_i, val in enumerate(row, 1):
            ws.cell(row=r_i, column=c_i, value=val)

    styler = ExcelStyler(theme="modern_slate", currency="IDR")
    styler.style_worksheet(ws, df, add_total_row=True)

    # Cek baris total dibuat di row 4
    assert ws.cell(row=4, column=1).value == "TOTAL"
    assert ws.cell(row=4, column=2).value == "=SUM(B2:B3)"

    # Cek auto-filter aktif
    assert ws.auto_filter.ref is not None
