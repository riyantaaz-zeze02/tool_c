"""
reporter.py — Generator Sheet Laporan Audit & Ringkasan Pembersihan (Cleaning Summary)
====================================================================================
Menghasilkan lembar kerja Excel kedua ("Cleaning Summary")
yang mencatat metrik transparansi data sebelum vs sesudah.
"""

from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from cleaner.formatter.themes import get_theme


def _translate_action(action_str: str) -> str:
    """Menerjemahkan deskripsi aksi pembersihan ke bahasa Inggris."""
    replacements = [
        ("strip spasi", "strip whitespace"),
        ("hapus non-ASCII", "remove non-ASCII"),
        ("persen → float (desimal)", "percentage → float (decimal)"),
        ("string angka → float", "numeric string → float"),
        ("(numerik)", "(numeric)"),
        ("(datetime)", "(datetime)"),
        ("(teks)", "(text)"),
        ("diisi rata-rata (mean)", "imputed with mean"),
        ("diisi median", "imputed with median"),
        ("diisi modus", "imputed with mode"),
        ("diisi konstanta", "filled with constant"),
        ("diisi 'Tidak Diketahui'", "filled with 'Unknown'"),
        ("diisi", "filled with"),
        ("dikonversi ke datetime", "converted to datetime"),
        ("(dibiarkan)", "(retained)"),
        ("Data sudah bersih", "Data already clean"),
    ]
    res = action_str
    for id_term, en_term in replacements:
        res = res.replace(id_term, en_term)
    return res


class CleaningReporter:
    """
    Membuat sheet audit trail dan ringkasan metrik pembersihan pada workbook openpyxl.
    """

    def __init__(self, theme="corporate_blue"):
        self.theme = get_theme(theme)

    def add_summary_sheet(self, wb, report_dict, df_cleaned=None, title="Cleaning Summary"):
        """
        Menambahkan worksheet 'Cleaning Summary' ke workbook Excel.
        Mendukung laporan single-sheet maupun multi-sheet (laporan terpisah per sheet).

        Parameter:
            wb (openpyxl.Workbook): Workbook yang sedang dibuat.
            report_dict (dict | list): Laporan dari CleaningEngine.clean() atau dict per sheet.
            df_cleaned (pd.DataFrame | dict | None): Dataframe hasil pembersihan.
            title (str): Judul sheet summary (default: 'Cleaning Summary').
        """
        ws = wb.create_sheet(title=title)
        ws.views.sheetView[0].showGridLines = True

        font_family = self.theme["font_name"]

        # Styles
        title_font = Font(name=font_family, size=14, bold=True, color=self.theme["header_fill"])
        subtitle_font = Font(name=font_family, size=9, italic=True, color="595959")
        section_font = Font(name=font_family, size=11, bold=True, color="262626")
        label_font = Font(name=font_family, size=10, bold=True, color="404040")
        value_font = Font(name=font_family, size=10)

        header_fill = PatternFill(
            start_color=self.theme["header_fill"],
            end_color=self.theme["header_fill"],
            fill_type="solid"
        )
        header_font = Font(name=font_family, size=10, bold=True, color=self.theme["header_font"])

        border_side = Side(border_style="thin", color=self.theme["border_color"])
        thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        zebra_fill = PatternFill(
            start_color=self.theme["zebra_fill"],
            end_color=self.theme["zebra_fill"],
            fill_type="solid"
        )

        # 1. Judul & Waktu (Header & Subtitle)
        ws.cell(row=2, column=2, value="DATA CLEANING AUDIT REPORT").font = title_font
        waktu_str = datetime.now().strftime("%d %B %Y, %H:%M:%S")
        ws.cell(row=3, column=2, value=f"YANTTT — {waktu_str}").font = subtitle_font

        # Normalisasi input ke list of (sheet_name, report, df)
        sheet_reports = []
        if isinstance(report_dict, dict) and ("langkah" in report_dict or "baris_awal" in report_dict):
            # Single sheet report
            s_name = report_dict.get("sumber", "Data")
            sheet_reports.append((s_name, report_dict, df_cleaned))
        elif isinstance(report_dict, dict):
            # Dict of sheets: {"Sheet1": {"report": ..., "df": ...}} or {"Sheet1": report_dict}
            for s_name, val in report_dict.items():
                if isinstance(val, dict) and "report" in val:
                    sheet_reports.append((s_name, val["report"], val.get("df")))
                elif isinstance(val, dict):
                    d = df_cleaned.get(s_name) if isinstance(df_cleaned, dict) else None
                    sheet_reports.append((s_name, val, d))
        elif isinstance(report_dict, list):
            for item in report_dict:
                if isinstance(item, tuple) and len(item) == 3:
                    sheet_reports.append(item)
                elif isinstance(item, tuple) and len(item) == 2:
                    sheet_reports.append((item[0], item[1], None))

        is_join = isinstance(report_dict, dict) and report_dict.get("is_join", False)
        is_merge = isinstance(report_dict, dict) and report_dict.get("is_merge", False)

        if is_join:
            ws.cell(row=5, column=2, value="JOIN FLOW").font = section_font
            metrics = [
                ("Mode Operasi", "Chained Left Join & Cleaning"),
                ("Total File Di-join", report_dict.get("total_files", 0)),
                ("Total Baris Awal", report_dict.get("baris_awal", 0)),
                ("Total Baris Akhir", report_dict.get("baris_akhir", 0)),
                ("Total Duplikat Dihapus", report_dict.get("total_duplikat_dihapus", 0)),
            ]
            for idx, (label, val) in enumerate(metrics, start=6):
                c_label = ws.cell(row=idx, column=2, value=label)
                c_label.font = label_font
                c_label.border = thin_border
                c_label.fill = zebra_fill
                c_val = ws.cell(row=idx, column=3, value=val)
                c_val.font = value_font
                c_val.border = thin_border
                if isinstance(val, (int, float)):
                    c_val.alignment = Alignment(horizontal="right")
                    c_val.number_format = "#,##0"

            row_stage_start = 13
            ws.cell(row=row_stage_start - 1, column=2, value="JOIN STAGES").font = section_font
            stage_headers = ["Stage", "Left Side", "Right File", "Key", "Matched Rows", "Unmatched Rows", "Match Rate"]
            for col_i, header in enumerate(stage_headers, start=2):
                cell = ws.cell(row=row_stage_start, column=col_i, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            for stage_index, stage in enumerate(report_dict.get("join_stages", []), start=1):
                row = row_stage_start + stage_index
                row_values = [
                    stage_index,
                    stage["left"],
                    stage["right"],
                    stage["key"],
                    stage["matched_rows"],
                    stage["unmatched_rows"],
                    f"{stage['match_rate']:.2f}%",
                ]
                for col_i, value in enumerate(row_values, start=2):
                    cell = ws.cell(row=row, column=col_i, value=value)
                    cell.font = value_font
                    cell.fill = zebra_fill if stage_index % 2 else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
                    cell.border = thin_border
                    if col_i in (2, 6, 7):
                        cell.alignment = Alignment(horizontal="right")

            row_flow = row_stage_start + len(report_dict.get("join_stages", [])) + 3
            ws.cell(row=row_flow - 1, column=2, value="JOIN FLOW SUMMARY").font = section_font
            for flow_index, stage in enumerate(report_dict.get("join_stages", []), start=row_flow):
                flow_text = f"{stage['left']} + {stage['right']} (key: {stage['key']}) → {stage['match_rate']:.2f}% match"
                ws.cell(row=flow_index, column=2, value=flow_text).font = value_font

            orphan_count = report_dict.get("total_orphan_rows", 0)
            orphan_note_row = row_flow + len(report_dict.get("join_stages", [])) + 1
            ws.cell(
                row=orphan_note_row,
                column=2,
                value=f"{orphan_count} baris tidak ikut ke hasil akhir — lihat sheet 'Baris Ter-drop' untuk detail",
            ).font = label_font

            for col_i in range(2, 9):
                col_letter = get_column_letter(col_i)
                max_len = max(
                    len(str(ws.cell(row=row, column=col_i).value or ""))
                    for row in range(2, orphan_note_row + 1)
                )
                ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        elif is_merge:
            # Layout khusus untuk laporan hasil penggabungan (Merge) (Requirement 6)
            ws.cell(row=5, column=2, value="KEY METRICS").font = section_font

            metrics = [
                ("Mode Operasi", "File Merge & Cleaning"),
                ("Total File Digabung", report_dict.get("total_files", 0)),
                ("Total Baris Awal", report_dict.get("baris_awal", 0)),
                ("Total Baris Akhir", report_dict.get("baris_akhir", 0)),
                ("Total Duplikat Dihapus", report_dict.get("total_duplikat_dihapus", 0)),
                ("  • Duplikat Internal (Satu File)", report_dict.get("duplikat_internal", 0)),
                ("  • Duplikat Antar-File (Cross-File)", report_dict.get("duplikat_antar_file", 0)),
                ("Baris Di-drop (Missing)", report_dict.get("total_baris_didrop", 0)),
                ("Total Kolom", report_dict.get("kolom_akhir", 0)),
            ]

            for idx, (label, val) in enumerate(metrics, start=6):
                c_label = ws.cell(row=idx, column=2, value=label)
                c_label.font = label_font
                c_label.border = thin_border
                c_label.fill = zebra_fill

                c_val = ws.cell(row=idx, column=3, value=val)
                c_val.font = value_font
                c_val.border = thin_border
                if isinstance(val, (int, float)):
                    c_val.alignment = Alignment(horizontal="right")
                    c_val.number_format = "#,##0"

            # Tabel Rincian File Sumber (Requirement 6)
            row_file_start = 6 + len(metrics) + 2
            ws.cell(row=row_file_start - 1, column=2, value="SOURCE FILES BREAKDOWN").font = section_font

            file_headers = ["No", "Source File Name", "Initial Rows", "Intra-file Duplicates"]
            for col_i, h_text in enumerate(file_headers, start=2):
                cell = ws.cell(row=row_file_start, column=col_i, value=h_text)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            file_details = report_dict.get("file_details", [])
            for f_idx, f_info in enumerate(file_details, start=1):
                r = row_file_start + f_idx
                is_z = (f_idx % 2 == 1)
                r_fill = zebra_fill if is_z else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

                f_row_data = [
                    f_idx,
                    f_info.get("file", ""),
                    f_info.get("baris_awal", 0),
                    f_info.get("duplikat_internal", 0),
                ]
                for c_i, val in enumerate(f_row_data, start=2):
                    cell = ws.cell(row=r, column=c_i, value=val)
                    cell.font = value_font
                    cell.fill = r_fill
                    cell.border = thin_border
                    if c_i == 2:
                        cell.alignment = Alignment(horizontal="center")
                    elif c_i >= 4:
                        cell.alignment = Alignment(horizontal="right")
                        cell.number_format = "#,##0"

            # Tabel Detail Status Kolom
            row_detail_start = row_file_start + len(file_details) + 2
            ws.cell(row=row_detail_start - 1, column=2, value="COLUMN STATUS DETAILS").font = section_font

            col_headers = ["No", "Column Name", "Final Data Type", "Missing Found", "Cleaning Action"]
            for col_i, h_text in enumerate(col_headers, start=2):
                cell = ws.cell(row=row_detail_start, column=col_i, value=h_text)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            if df_cleaned is not None:
                missing_log = report_dict.get("langkah", {}).get("missing", {}).get("missing_sebelum", {})
                actions_list = []
                for step in ["number", "date", "missing", "text"]:
                    step_actions = report_dict.get("langkah", {}).get(step, {}).get("aksi", [])
                    actions_list.extend(step_actions)

                for col_idx, col_name in enumerate(df_cleaned.columns, start=1):
                    r = row_detail_start + col_idx
                    is_z = (col_idx % 2 == 1)
                    r_fill = zebra_fill if is_z else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

                    col_actions = [a for a in actions_list if f"'{col_name}'" in a]
                    if col_actions:
                        translated_actions = [_translate_action(a) for a in col_actions]
                        action_desc = "; ".join(translated_actions)
                    else:
                        action_desc = "Data already clean"

                    col_row_data = [
                        col_idx,
                        col_name,
                        str(df_cleaned[col_name].dtype),
                        missing_log.get(col_name, 0),
                        action_desc,
                    ]
                    for c_i, val in enumerate(col_row_data, start=2):
                        cell = ws.cell(row=r, column=c_i, value=val)
                        cell.font = value_font
                        cell.fill = r_fill
                        cell.border = thin_border
                        if c_i == 2:
                            cell.alignment = Alignment(horizontal="center")
                        elif c_i == 5:
                            cell.alignment = Alignment(horizontal="right")
                            cell.number_format = "#,##0"

            # Auto-width kolom
            total_rows_to_check = row_detail_start + (len(df_cleaned.columns) if df_cleaned is not None else 6) + 2
            for c in range(2, 7):
                c_letter = get_column_letter(c)
                max_len = 0
                for r in range(2, total_rows_to_check):
                    v = ws.cell(row=r, column=c).value
                    if v is not None:
                        max_len = max(max_len, len(str(v)))
                ws.column_dimensions[c_letter].width = max(max_len + 4, 15)

        # Jika hanya 1 sheet: gunakan format klasik yang kompatibel penuh dengan test yang ada
        elif len(sheet_reports) <= 1:
            single_rep = sheet_reports[0][1] if sheet_reports else (report_dict if isinstance(report_dict, dict) else {})
            single_df = sheet_reports[0][2] if sheet_reports else df_cleaned

            # 2. Ringkasan Metrik Kunci (Key Metrics)
            ws.cell(row=5, column=2, value="KEY METRICS").font = section_font


            metrics = [
                ("Source File", single_rep.get("sumber", "Input File")),
                ("Initial Row Count", single_rep.get("baris_awal", 0)),
                ("Final Row Count", single_rep.get("baris_akhir", 0)),
                ("Duplicate Rows Removed", single_rep.get("total_duplikat_dihapus", 0)),
                ("Dropped Rows (Missing)", single_rep.get("total_baris_didrop", 0)),
                ("Total Columns", single_rep.get("kolom_akhir", 0)),
            ]

            for idx, (label, val) in enumerate(metrics, start=6):
                c_label = ws.cell(row=idx, column=2, value=label)
                c_label.font = label_font
                c_label.border = thin_border
                c_label.fill = zebra_fill

                c_val = ws.cell(row=idx, column=3, value=val)
                c_val.font = value_font
                c_val.border = thin_border
                if isinstance(val, (int, float)):
                    c_val.alignment = Alignment(horizontal="right")
                    c_val.number_format = "#,##0"

            # 3. Tabel Detail Perubahan per Kolom (Column Status Details)
            row_detail_start = 14
            ws.cell(row=row_detail_start - 1, column=2, value="COLUMN STATUS DETAILS").font = section_font

            headers = ["No", "Column Name", "Final Data Type", "Missing Found", "Cleaning Action"]
            for col_i, h_text in enumerate(headers, start=2):
                cell = ws.cell(row=row_detail_start, column=col_i, value=h_text)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            if single_df is not None:
                missing_log = single_rep.get("langkah", {}).get("missing", {}).get("missing_sebelum", {})
                actions_list = []
                for step in ["number", "date", "missing", "text"]:
                    step_actions = single_rep.get("langkah", {}).get(step, {}).get("aksi", [])
                    actions_list.extend(step_actions)

                for col_idx, col_name in enumerate(single_df.columns, start=1):
                    r = row_detail_start + col_idx
                    is_z = (col_idx % 2 == 1)
                    row_fill = zebra_fill if is_z else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

                    col_actions = [a for a in actions_list if f"'{col_name}'" in a]
                    if col_actions:
                        translated_actions = [_translate_action(a) for a in col_actions]
                        action_desc = "; ".join(translated_actions)
                    else:
                        action_desc = "Data already clean"

                    row_data = [
                        col_idx,
                        col_name,
                        str(single_df[col_name].dtype),
                        missing_log.get(col_name, 0),
                        action_desc,
                    ]

                    for c_i, val in enumerate(row_data, start=2):
                        cell = ws.cell(row=r, column=c_i, value=val)
                        cell.font = value_font
                        cell.fill = row_fill
                        cell.border = thin_border
                        if c_i == 2:
                            cell.alignment = Alignment(horizontal="center")
                        elif c_i == 5:
                            cell.alignment = Alignment(horizontal="right")
                            cell.number_format = "#,##0"

            # Auto-width kolom
            for c in range(2, 7):
                c_letter = get_column_letter(c)
                max_len = 0
                max_row = row_detail_start + (len(single_df.columns) if single_df is not None else 6) + 2
                for r in range(2, max_row):
                    v = ws.cell(row=r, column=c).value
                    if v is not None:
                        max_len = max(max_len, len(str(v)))
                ws.column_dimensions[c_letter].width = max(max_len + 4, 15)

        else:
            # Multi-sheet: Laporan TERPISAH per sheet (Requirement 5)
            # 2. Tabel Ringkasan Metrik Kunci per Sheet
            ws.cell(row=5, column=2, value="KEY METRICS").font = section_font

            table_headers = [
                "No", "Sheet Name", "Initial Rows", "Final Rows",
                "Duplicates Removed", "Dropped Rows", "Total Columns"
            ]
            for col_i, h_text in enumerate(table_headers, start=2):
                cell = ws.cell(row=6, column=col_i, value=h_text)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            current_row = 7
            for s_idx, (s_name, s_rep, s_df) in enumerate(sheet_reports, start=1):
                is_z = (s_idx % 2 == 1)
                row_fill = zebra_fill if is_z else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

                row_metrics = [
                    s_idx,
                    s_name,
                    s_rep.get("baris_awal", len(s_df) if s_df is not None else 0),
                    s_rep.get("baris_akhir", len(s_df) if s_df is not None else 0),
                    s_rep.get("total_duplikat_dihapus", 0),
                    s_rep.get("total_baris_didrop", 0),
                    s_rep.get("kolom_akhir", len(s_df.columns) if s_df is not None else 0),
                ]

                for col_i, val in enumerate(row_metrics, start=2):
                    cell = ws.cell(row=current_row, column=col_i, value=val)
                    cell.font = value_font
                    cell.fill = row_fill
                    cell.border = thin_border
                    if col_i == 2:
                        cell.alignment = Alignment(horizontal="center")
                    elif col_i >= 4:
                        cell.alignment = Alignment(horizontal="right")
                        cell.number_format = "#,##0"

                current_row += 1

            # 3. Rincian Kolom Terpisah untuk Setiap Sheet
            current_row += 2
            col_headers = ["No", "Column Name", "Final Data Type", "Missing Found", "Cleaning Action"]

            for s_name, s_rep, s_df in sheet_reports:
                ws.cell(row=current_row, column=2, value=f"COLUMN STATUS DETAILS — {s_name}").font = section_font
                current_row += 1

                for col_i, h_text in enumerate(col_headers, start=2):
                    cell = ws.cell(row=current_row, column=col_i, value=h_text)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thin_border
                current_row += 1

                if s_df is not None:
                    missing_log = s_rep.get("langkah", {}).get("missing", {}).get("missing_sebelum", {})
                    actions_list = []
                    for step in ["number", "date", "missing", "text"]:
                        step_actions = s_rep.get("langkah", {}).get(step, {}).get("aksi", [])
                        actions_list.extend(step_actions)

                    for col_idx, col_name in enumerate(s_df.columns, start=1):
                        is_z = (col_idx % 2 == 1)
                        row_fill = zebra_fill if is_z else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

                        col_actions = [a for a in actions_list if f"'{col_name}'" in a]
                        if col_actions:
                            translated_actions = [_translate_action(a) for a in col_actions]
                            action_desc = "; ".join(translated_actions)
                        else:
                            action_desc = "Data already clean"

                        row_data = [
                            col_idx,
                            col_name,
                            str(s_df[col_name].dtype),
                            missing_log.get(col_name, 0),
                            action_desc,
                        ]

                        for c_i, val in enumerate(row_data, start=2):
                            cell = ws.cell(row=current_row, column=c_i, value=val)
                            cell.font = value_font
                            cell.fill = row_fill
                            cell.border = thin_border
                            if c_i == 2:
                                cell.alignment = Alignment(horizontal="center")
                            elif c_i == 5:
                                cell.alignment = Alignment(horizontal="right")
                                cell.number_format = "#,##0"

                        current_row += 1

                current_row += 2

            # Auto-width kolom untuk semua baris
            for c in range(2, 9):
                c_letter = get_column_letter(c)
                max_len = 0
                for r in range(2, current_row):
                    v = ws.cell(row=r, column=c).value
                    if v is not None:
                        max_len = max(max_len, len(str(v)))
                ws.column_dimensions[c_letter].width = max(max_len + 4, 15)

