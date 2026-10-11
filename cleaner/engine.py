"""
engine.py — Pipeline Orkestrator Data Cleaning
===============================================
Mengatur urutan pembersihan data, mengelola konfigurasi,
dan mencatat audit trail (laporan statistik pembersihan).
"""

import pandas as pd
import os
import re
from cleaner.reader import baca_data, get_sheet_names
from cleaner.rules.missing_handler import tangani_missing

from cleaner.rules.duplicate_handler import hapus_duplikat
from cleaner.rules.text_cleaner import bersihkan_teks
from cleaner.rules.number_cleaner import bersihkan_angka
from cleaner.rules.date_cleaner import bersihkan_tanggal
from cleaner.formatter.excel_styler import ExcelStyler
from cleaner.reporter import CleaningReporter
import openpyxl


class CleaningEngine:
    """
    Orkestrator pipeline pembersihan data.
    """

    DEFAULT_CONFIG = {
        "missing": {
            "numeric_strategy": "median",
            "text_strategy": "constant",
            "numeric_constant": 0,
            "text_constant": "Unknown",
            "drop_threshold": None,
        },
        "duplicate": {
            "subset": None,
            "keep": "first",
        },
        "number": {
            "kolom_angka": [],
            "kolom_currency": [],
            "kolom_persen": [],
            "auto_detect": True,
        },
        "date": {
            "kolom_tanggal": [],
            "auto_detect": True,
            "output_format": None,
        },
        "text": {
            "case_format": "title",
            "only_name_columns": True,
            "kolom_case": [],
            "preserve_brackets": True,
            "preserve_acronyms": True,
            "custom_acronyms": [],
            "strip_whitespace": True,
            "normalize_unicode": True,
            "remove_non_ascii": False,
            "kolom_exclude": ["email", "e-mail"],
        },
    }

    def __init__(self, config=None):
        """
        Inisialisasi CleaningEngine dengan konfigurasi kustom atau default.
        """
        self.config = {}
        # Merge default config dengan custom config
        for step, defaults in self.DEFAULT_CONFIG.items():
            self.config[step] = defaults.copy()
            if config and step in config:
                self.config[step].update(config[step])

        self.last_report = {}

    def clean(self, data_or_path, sheet_name=0):
        """
        Menjalankan seluruh pipeline pembersihan pada file atau DataFrame.

        Parameter:
            data_or_path (str | pd.DataFrame): Path file atau objek DataFrame.
            sheet_name (str | int): Sheet Excel yang dibaca jika berupa file Excel.

        Return:
            tuple: (pd.DataFrame bersih, dict laporan_audit)
        """
        source_name = "DataFrame"
        if isinstance(data_or_path, str):
            source_name = os.path.basename(data_or_path)
            df = baca_data(data_or_path, sheet_name=sheet_name)
        elif isinstance(data_or_path, pd.DataFrame):
            df = data_or_path.copy()
        else:
            raise TypeError("data_or_path harus berupa string path file atau pd.DataFrame")

        if df is None:
            raise ValueError("Gagal membaca data.")

        initial_rows = len(df)
        initial_cols = len(df.columns)

        report = {
            "sumber": source_name,
            "baris_awal": initial_rows,
            "kolom_awal": initial_cols,
            "langkah": {},
        }

        # 1. Parsing Angka & Mata Uang (sebelum teks distrip/dimodifikasi)
        df, log_number = bersihkan_angka(df, config=self.config.get("number", {}))
        report["langkah"]["number"] = log_number

        # 2. Parsing Tanggal
        df, log_date = bersihkan_tanggal(df, config=self.config.get("date", {}))
        report["langkah"]["date"] = log_date

        # 3. Tangani Missing Value
        df, log_missing = tangani_missing(df, config=self.config.get("missing", {}))
        report["langkah"]["missing"] = log_missing

        # 4. Hapus Duplikat
        df, log_duplicate = hapus_duplikat(df, config=self.config.get("duplicate", {}))
        report["langkah"]["duplicate"] = log_duplicate

        # 5. Standarisasi Teks
        df, log_text = bersihkan_teks(df, config=self.config.get("text", {}))
        report["langkah"]["text"] = log_text

        # Ringkasan Akhir
        report["baris_akhir"] = len(df)
        report["kolom_akhir"] = len(df.columns)
        report["total_duplikat_dihapus"] = log_duplicate.get("jumlah_duplikat", 0)
        report["total_baris_didrop"] = log_missing.get("baris_didrop", 0)

        self.last_report = report
        return df, report

    def _drop_total_rows(self, df):
        """Hapus baris TOTAL yang sering ditambah export_excel di akhir file."""
        if df.empty or len(df.columns) == 0:
            return df.copy()
        first_col = df.columns[0]
        total_mask = df[first_col].astype(str).str.strip().str.upper().eq("TOTAL")
        if not total_mask.any():
            return df.copy()
        remaining_mask = df.drop(columns=[first_col]).isna().all(axis=1)
        mask = total_mask & remaining_mask
        if mask.any():
            df = df.loc[~mask].copy().reset_index(drop=True)
        return df

    def _coerce_for_compare(self, value):
        """Normalisasi tipe agar dedupe/append dibuat pada nilai yang setara."""
        if pd.isna(value):
            return None
        if isinstance(value, str):
            value = value.strip()
            if value == "":
                return None
            if re.fullmatch(r"\d{4}-\d{1,2}-\d{1,2}", value):
                try:
                    return pd.to_datetime(value).strftime("%Y-%m-%d")
                except Exception:
                    pass
            if re.fullmatch(r"\d{1,2}[/-]\d{1,2}[/-]\d{4}", value):
                try:
                    return pd.to_datetime(value, errors="coerce").strftime("%Y-%m-%d")
                except Exception:
                    pass
            if re.search(r"\d", value):
                numeric_only = re.sub(r"[^0-9,\.\-]", "", value)
                if numeric_only:
                    if "," in numeric_only and "." in numeric_only:
                        if numeric_only.rfind(",") > numeric_only.rfind("."):
                            numeric_only = numeric_only.replace(".", "").replace(",", ".")
                        else:
                            numeric_only = numeric_only.replace(",", "")
                    elif "," in numeric_only and "." not in numeric_only:
                        if re.match(r"^\d{1,3}(?:,\d{3})+(?:,\d+)?$", numeric_only):
                            numeric_only = numeric_only.replace(",", "")
                        else:
                            numeric_only = numeric_only.replace(",", ".")
                    elif "." in numeric_only and "," not in numeric_only:
                        if re.match(r"^\d{1,3}(?:\.\d{3})+$", numeric_only):
                            numeric_only = numeric_only.replace(".", "")
                    try:
                        return float(numeric_only)
                    except ValueError:
                        pass
            return value.lower()
        if isinstance(value, (int, float)):
            if pd.isna(value):
                return None
            return float(value)
        if hasattr(value, "to_pydatetime"):
            try:
                return value.to_pydatetime().strftime("%Y-%m-%d")
            except Exception:
                return str(value)
        return str(value).strip().lower()

    def _row_signature(self, row, dedupe_keys=None):
        """Buat kunci dedupe yang netral terhadap urutan kolom dan tipe dasar."""
        if dedupe_keys is None:
            keys = list(row.index)
        else:
            keys = [key for key in dedupe_keys if key in row.index]
        if not keys:
            keys = list(row.index)
        values = []
        for key in keys:
            values.append(self._coerce_for_compare(row.get(key)))
        return tuple(values)

    def append_files(self, base_path, new_paths, dedupe_keys=None, sheet_name=0, source_col=None):
        """Tambahkan baris dari file baru ke file dasar tanpa membersihkan ulang base."""
        if not base_path or not os.path.exists(base_path):
            raise FileNotFoundError(f"File dasar '{base_path}' tidak ditemukan!")
        if not new_paths:
            raise ValueError("Perlu minimal 1 file baru untuk operasi append.")
        if isinstance(new_paths, (str, os.PathLike)):
            new_paths = [str(new_paths)]
        if len(new_paths) == 0:
            raise ValueError("Perlu minimal 1 file baru untuk operasi append.")

        base_df = baca_data(base_path, sheet_name=sheet_name).copy()
        base_rows_before = len(base_df)
        base_df = self._drop_total_rows(base_df)
        base_total_row_dibuang = base_rows_before - len(base_df)

        if source_col is not None:
            source_col = str(source_col)

        base_cols = list(base_df.columns)
        if source_col and source_col not in base_cols:
            base_df[source_col] = os.path.basename(base_path)

        base_seen = set()
        for _, row in base_df.iterrows():
            key = self._row_signature(row, dedupe_keys=dedupe_keys or base_cols)
            base_seen.add(key)

        final_frames = [base_df.copy()]
        skipped_rows = []
        file_details = []
        total_added = 0
        total_skipped = 0
        langkah = []

        for new_path in new_paths:
            if not os.path.exists(new_path):
                raise FileNotFoundError(f"File baru '{new_path}' tidak ditemukan!")
            raw_new = baca_data(new_path, sheet_name=sheet_name).copy()
            current_name = os.path.basename(new_path)
            raw_new = self._drop_total_rows(raw_new)

            if set(map(str, raw_new.columns)) != set(map(str, base_cols)):
                base_names = [str(c) for c in base_cols]
                incoming_names = [str(c) for c in raw_new.columns]
                missing = [c for c in base_names if c not in incoming_names]
                extra = [c for c in incoming_names if c not in base_names]
                err_lines = [
                    f"Struktur kolom tidak sama! File '{current_name}' berbeda dengan file acuan '{os.path.basename(base_path)}'.",
                    f"   • Kolom acuan ({os.path.basename(base_path)}): {base_cols}",
                    f"   • Kolom ditemukan ({current_name}): {list(raw_new.columns)}",
                ]
                if missing:
                    err_lines.append(f"   • Kolom hilang di '{current_name}': {missing}")
                if extra:
                    err_lines.append(f"   • Kolom tambahan di '{current_name}': {extra}")
                raise ValueError("\n".join(err_lines))

            new_df = raw_new.reindex(columns=base_cols).copy()
            if source_col is not None:
                new_df[source_col] = current_name
            cleaned_new, clean_report = self.clean(new_df)
            langkah.append({"file": current_name, "report": clean_report})

            seen_in_file = set()
            kept_rows = []
            kept_keys = []
            for _, row in cleaned_new.iterrows():
                key = self._row_signature(row, dedupe_keys=dedupe_keys or list(cleaned_new.columns))
                if key in base_seen:
                    skipped_rows.append({**row.to_dict(), "File Asal": current_name, "Alasan": "sudah ada di file dasar"})
                    total_skipped += 1
                    continue
                if key in seen_in_file:
                    skipped_rows.append({**row.to_dict(), "File Asal": current_name, "Alasan": "duplikat di file baru"})
                    total_skipped += 1
                    continue
                seen_in_file.add(key)
                kept_rows.append(row.to_dict())
                kept_keys.append(key)
                total_added += 1
            base_seen.update(kept_keys)

            added_df = pd.DataFrame(kept_rows, columns=cleaned_new.columns)
            if not added_df.empty:
                final_frames.append(added_df)

            file_details.append({
                "file": current_name,
                "baris_awal": len(raw_new),
                "ditambahkan": len(added_df),
                "dilewati_vs_base": sum(1 for row in skipped_rows if row.get("File Asal") == current_name and row.get("Alasan") == "sudah ada di file dasar"),
                "dilewati_dalam_file_baru": sum(1 for row in skipped_rows if row.get("File Asal") == current_name and row.get("Alasan") == "duplikat di file baru"),
            })

        final_df = pd.concat(final_frames, ignore_index=True)

        if len(final_df) > 1_048_576:
            raise ValueError("Data terlalu besar untuk satu sheet Excel (maks 1.048.576 baris)")

        skipped_df = pd.DataFrame(skipped_rows)
        if skipped_df.empty:
            skipped_df = pd.DataFrame(columns=list(base_df.columns) + ["File Asal", "Alasan"])

        report = {
            "sumber": f"Append ({len(new_paths)} files)",
            "is_append": True,
            "source_col": source_col,
            "base_rows": len(base_df),
            "base_total_row_dibuang": base_total_row_dibuang,
            "file_details": file_details,
            "total_ditambahkan": total_added,
            "total_dilewati": total_skipped,
            "baris_awal": len(base_df) + sum(len(pd.read_csv(p)) if os.path.splitext(p)[1].lower() == ".csv" else len(pd.read_excel(p)) for p in new_paths),
            "baris_akhir": len(final_df),
            "kolom_awal": len(base_df.columns),
            "kolom_akhir": len(final_df.columns),
            "langkah": langkah,
            "skipped_rows": skipped_df,
        }
        self.last_report = report
        return final_df, report

    def clean_sheets(self, file_path, sheets=None, all_sheets=False):
        """
        Membersihkan satu atau beberapa sheet dari file Excel/CSV.

        Parameter:
            file_path (str): Path ke file input.
            sheets (list[str|int] | str | None): Sheet yang ingin dibersihkan.
            all_sheets (bool): Jika True, bersihkan semua sheet yang ada di file.

        Return:
            dict: {
                sheet_name: {
                    "df": pd.DataFrame,
                    "report": dict,
                }
            }
        """
        if not isinstance(file_path, str):
            raise TypeError("file_path harus berupa string path file")

        ekstensi = os.path.splitext(file_path)[1].lower()
        if ekstensi in [".xlsx", ".xls"]:
            available_sheets = get_sheet_names(file_path)
            if all_sheets:
                target_sheets = available_sheets
            elif sheets is not None:
                if isinstance(sheets, str):
                    target_sheets = [s.strip() for s in sheets.split(",") if s.strip()]
                elif isinstance(sheets, (list, tuple)):
                    target_sheets = list(sheets)
                else:
                    target_sheets = [sheets]
            else:
                target_sheets = [available_sheets[0]] if available_sheets else [0]

            # Validasi nama sheet
            for s in target_sheets:
                if isinstance(s, str) and s not in available_sheets:
                    raise ValueError(
                        f"Sheet '{s}' tidak ditemukan di file. Sheet yang tersedia: {', '.join(available_sheets)}"
                    )

            cleaned_results = {}
            for s in target_sheets:
                df_clean, rep = self.clean(file_path, sheet_name=s)
                cleaned_results[str(s)] = {"df": df_clean, "report": rep}
            return cleaned_results
        else:
            # File CSV atau single table
            df_clean, rep = self.clean(file_path)
            s_name = os.path.splitext(os.path.basename(file_path))[0]
            return {s_name: {"df": df_clean, "report": rep}}

    def clean_merged_files(self, file_paths, source_col="Sumber File", sheet_name=0):
        """
        Menggabungkan beberapa file dengan struktur kolom yang sama,
        menambahkan kolom asal file, dan membersihkan seluruh dataset secara terpadu.

        Parameter:
            file_paths (list[str]): Daftar path file yang akan digabungkan.
            source_col (str): Nama kolom untuk mencatat file sumber asal (default: 'Sumber File').
            sheet_name (str|int): Sheet yang dibaca jika file berupa Excel.

        Return:
            tuple: (pd.DataFrame bersih hasil gabungan, dict laporan_audit)
        """
        if not file_paths or len(file_paths) < 2:
            raise ValueError("Perlu minimal 2 file untuk operasi penggabungan (merge).")

        raw_files = self.validate_merged_files(file_paths, sheet_name=sheet_name)

        # 1. Baca semua file dan catat metrik awal
        file_list = [fname for _, fname, _ in raw_files]

        # 2. Tandai asal file di setiap baris dan tumpuk dengan pd.concat
        tagged_dfs = []
        for _, fname, df in raw_files:
            df_tagged = df.copy()
            df_tagged[source_col] = fname
            tagged_dfs.append(df_tagged)

        df_merged = pd.concat(tagged_dfs, ignore_index=True)
        total_initial_rows = len(df_merged)

        # 3. Deteksi duplikat intra-file vs antar-file sebelum deduplikasi
        feat_cols = [c for c in df_merged.columns if c != source_col]
        first_seen_file = {}
        intra_dups_per_file = {fname: 0 for fname in file_list}
        intra_dups_total = 0
        cross_dups_total = 0

        for _, row in df_merged.iterrows():
            row_key = tuple(row[feat_cols].astype(str))
            file_src = row[source_col]
            if row_key not in first_seen_file:
                first_seen_file[row_key] = file_src
            else:
                if first_seen_file[row_key] == file_src:
                    intra_dups_per_file[file_src] += 1
                    intra_dups_total += 1
                else:
                    cross_dups_total += 1

        # 4. Jalankan cleaning pipeline pada hasil gabungan
        dup_config = self.config.get("duplicate", {}).copy()
        if dup_config.get("subset") is None:
            dup_config["subset"] = feat_cols

        text_config = self.config.get("text", {}).copy()
        excludes = text_config.get("kolom_exclude", [])
        if source_col not in excludes:
            text_config["kolom_exclude"] = list(excludes) + [source_col]

        report = {
            "sumber": f"Merged ({len(file_paths)} files)",
            "is_merge": True,
            "source_col": source_col,
            "total_files": len(file_paths),
            "file_details": [
                {
                    "file": fname,
                    "baris_awal": next(len(df) for _, current_name, df in raw_files if current_name == fname),
                    "duplikat_internal": intra_dups_per_file[fname],
                }
                for fname in file_list
            ],
            "baris_awal": total_initial_rows,
            "duplikat_internal": intra_dups_total,
            "duplikat_antar_file": cross_dups_total,
            "kolom_awal": len(df_merged.columns),
            "langkah": {},
        }

        # Step 1: Angka & Currency
        df_merged, log_number = bersihkan_angka(df_merged, config=self.config.get("number", {}))
        report["langkah"]["number"] = log_number

        # Step 2: Parsing Tanggal
        df_merged, log_date = bersihkan_tanggal(df_merged, config=self.config.get("date", {}))
        report["langkah"]["date"] = log_date

        # Step 3: Missing Value
        df_merged, log_missing = tangani_missing(df_merged, config=self.config.get("missing", {}))
        report["langkah"]["missing"] = log_missing

        # Step 4: Deduplikasi
        df_merged, log_duplicate = hapus_duplikat(df_merged, config=dup_config)
        report["langkah"]["duplicate"] = log_duplicate

        # Step 5: Standarisasi Teks
        df_merged, log_text = bersihkan_teks(df_merged, config=text_config)
        report["langkah"]["text"] = log_text

        report["baris_akhir"] = len(df_merged)
        report["kolom_akhir"] = len(df_merged.columns)
        report["total_duplikat_dihapus"] = log_duplicate.get("jumlah_duplikat", 0)
        report["total_baris_didrop"] = log_missing.get("baris_didrop", 0)

        self.last_report = report
        return df_merged, report

    def validate_merged_files(self, file_paths, sheet_name=0):
        """Membaca file merge dan memvalidasi urutan kolomnya."""
        if not file_paths or len(file_paths) < 2:
            raise ValueError("Perlu minimal 2 file untuk operasi penggabungan (merge).")

        raw_files = []
        for f in file_paths:
            if not os.path.exists(f):
                raise FileNotFoundError(f"File '{f}' tidak ditemukan!")
            df = baca_data(f, sheet_name=sheet_name)
            fname = os.path.basename(f)
            raw_files.append((f, fname, df.copy()))

        file_list = [fname for _, fname, _ in raw_files]
        base_file = file_list[0]
        base_cols = list(raw_files[0][2].columns)

        for _, curr_file, curr_df in raw_files[1:]:
            curr_cols = list(curr_df.columns)
            if curr_cols != base_cols:
                missing_in_curr = [c for c in base_cols if c not in curr_cols]
                extra_in_curr = [c for c in curr_cols if c not in base_cols]
                err_lines = [
                    f"Struktur kolom tidak sama! File '{curr_file}' berbeda dengan file acuan '{base_file}'.",
                    f"   • Kolom acuan ({base_file}): {base_cols}",
                    f"   • Kolom ditemukan ({curr_file}): {curr_cols}",
                ]
                if missing_in_curr:
                    err_lines.append(f"   • Kolom hilang di '{curr_file}': {missing_in_curr}")
                if extra_in_curr:
                    err_lines.append(f"   • Kolom tambahan di '{curr_file}': {extra_in_curr}")
                raise ValueError("\n".join(err_lines))
        return raw_files

    def clean_joined_files(self, file_paths, keys, sheet_name=0, how="left"):
        """Melakukan chained join pada beberapa file lalu membersihkan hasil akhirnya."""
        if not file_paths or len(file_paths) < 2:
            raise ValueError("Perlu minimal 2 file untuk operasi join.")
        if isinstance(keys, str):
            keys = [key.strip() for key in keys.split(",") if key.strip()]
        else:
            keys = [str(key).strip() for key in keys if str(key).strip()]
        if len(keys) != len(file_paths) - 1:
            raise ValueError(
                f"Jumlah key tidak sesuai: {len(file_paths)} file membutuhkan "
                f"{len(file_paths) - 1} key, tetapi menerima {len(keys)}."
            )
        if how != "left":
            raise ValueError("Chained join hanya mendukung tipe join 'left'.")

        raw_files = []
        for file_path in file_paths:
            if not os.path.exists(file_path):
                raise FileNotFoundError(f"File '{file_path}' tidak ditemukan!")
            raw_files.append((os.path.basename(file_path), baca_data(file_path, sheet_name=sheet_name).copy()))

        # Validasi seluruh rantai sebelum satu pun join dijalankan.
        result_columns = list(raw_files[0][1].columns)
        for stage_index, key in enumerate(keys):
            left_name = "hasil gabungan" if stage_index else raw_files[0][0]
            right_name, right_df = raw_files[stage_index + 1]
            if key not in result_columns:
                raise ValueError(
                    f"Key '{key}' untuk tahap {stage_index + 1} tidak ditemukan "
                    f"di {left_name}. Kolom tersedia: {result_columns}"
                )
            if key not in right_df.columns:
                raise ValueError(
                    f"Key '{key}' untuk tahap {stage_index + 1} tidak ditemukan "
                    f"di file '{right_name}'. Kolom tersedia: {list(right_df.columns)}"
                )
            left_empty = pd.DataFrame(columns=result_columns)
            right_empty = pd.DataFrame(columns=list(right_df.columns))
            result_columns = list(left_empty.merge(right_empty, how=how, on=key).columns)

        joined = raw_files[0][1]
        join_stages = []
        orphan_rows = []
        for stage_index, key in enumerate(keys):
            right_name, right_df = raw_files[stage_index + 1]
            left_name = "hasil gabungan" if stage_index else raw_files[0][0]
            left_keys = joined[key].dropna().unique()
            right_keys = right_df[key].dropna().unique()
            orphan_mask = ~right_df[key].isin(left_keys)
            for _, orphan_row in right_df.loc[orphan_mask].iterrows():
                orphan_data = orphan_row.to_dict()
                orphan_value = orphan_row[key]
                if pd.isna(orphan_value):
                    orphan_value = "<kosong>"
                orphan_data.update({
                    "Tahap": stage_index + 1,
                    "File Asal": right_name,
                    "Alasan": f"{key} {orphan_value} tidak ditemukan di {left_name}",
                })
                orphan_rows.append(orphan_data)

            matched_mask = joined[key].isin(right_keys)
            matched_rows = int(matched_mask.sum())
            left_rows = len(joined)
            stage = {
                "tahap": stage_index + 1,
                "left": left_name,
                "right": right_name,
                "key": key,
                "left_rows": left_rows,
                "matched_rows": matched_rows,
                "unmatched_rows": left_rows - matched_rows,
                "match_rate": round((matched_rows / left_rows * 100) if left_rows else 0, 2),
            }
            join_stages.append(stage)
            joined = joined.merge(right_df, how=how, on=key, suffixes=("", f"_{os.path.splitext(right_name)[0]}"))

        initial_rows = len(joined)
        cleaned, report = self.clean(joined)
        report.update({
            "sumber": f"Joined ({len(file_paths)} files)",
            "is_join": True,
            "total_files": len(file_paths),
            "join_keys": keys,
            "join_stages": join_stages,
            "orphan_rows": pd.DataFrame(orphan_rows),
            "total_orphan_rows": len(orphan_rows),
            "baris_awal": initial_rows,
            "kolom_awal": len(joined.columns),
        })
        self.last_report = report
        return cleaned, report

    def export_excel(

        self,
        df,
        output_path,
        report=None,
        theme="corporate_blue",
        currency="IDR",
        sheet_name="Data Bersih",
        add_total_row=True,
        include_summary_sheet=True,
        original_file_path=None,
    ):
        """
        Mengekspor DataFrame atau multi-sheet hasil pembersihan ke file Excel dengan styling profesional.
        Mendukung pemeliharaan sheet unselected dari original_file_path secara utuh (Requirement 4).

        Parameter:
            df (pd.DataFrame | dict): DataFrame bersih atau dict {sheet_name: df} atau {sheet_name: {"df": df, "report": ...}}.
            output_path (str): Lokasi file output (.xlsx).
            report (dict|None): Dict audit report (atau dict multi-sheet report).
            theme (str): Tema warna Excel ('corporate_blue', 'modern_slate', 'emerald_green', dll).
            currency (str): Mata uang ('IDR', 'USD').
            sheet_name (str): Nama worksheet data utama (jika df berupa DataFrame tunggal).
            add_total_row (bool): Apakah menyertakan baris Total dengan formula =SUM.
            include_summary_sheet (bool): Apakah menyertakan sheet 'Cleaning Summary'.
            original_file_path (str|None): Path ke file Excel sumber asli untuk menyalin sheet unselected apa adanya.
        """
        # Normalisasi struktur input dict vs single DataFrame
        is_multi = isinstance(df, dict)
        if is_multi:
            sheet_dfs = {}
            sheet_reports = {}
            for k, v in df.items():
                if isinstance(v, dict) and "df" in v:
                    sheet_dfs[str(k)] = v["df"]
                    if "report" in v:
                        sheet_reports[str(k)] = v["report"]
                else:
                    sheet_dfs[str(k)] = v
            if isinstance(report, dict):
                sheet_reports.update(report)
        else:
            sheet_dfs = {sheet_name: df}
            sheet_reports = {sheet_name: report} if report else {}

        # Validasi batas baris maksimum Excel (.xlsx) untuk semua sheet
        for s_n, s_df in sheet_dfs.items():
            if len(s_df) > 1_048_576:
                raise ValueError("Data terlalu besar untuk satu sheet Excel (maks 1.048.576 baris)")

        active_report_for_extras = report if isinstance(report, dict) else self.last_report
        orphan_df = None
        skipped_append_df = None
        if isinstance(active_report_for_extras, dict) and active_report_for_extras.get("is_join"):
            orphan_df = active_report_for_extras.get("orphan_rows")
            if orphan_df is None or orphan_df.empty:
                orphan_df = pd.DataFrame(columns=["Tahap", "File Asal", "Alasan"])
            if len(orphan_df) > 1_048_576:
                raise ValueError("Data orphan terlalu besar untuk satu sheet Excel (maks 1.048.576 baris)")
        if isinstance(active_report_for_extras, dict) and active_report_for_extras.get("is_append"):
            skipped_append_df = active_report_for_extras.get("skipped_rows")
            if skipped_append_df is None or skipped_append_df.empty:
                skipped_append_df = pd.DataFrame(columns=["File Asal", "Alasan"])
            if len(skipped_append_df) > 1_048_576:
                raise ValueError("Data baris dilewati terlalu besar untuk satu sheet Excel (maks 1.048.576 baris)")

        # Pastikan direktori output ada
        out_dir = os.path.dirname(output_path)
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)

        styler = ExcelStyler(theme=theme, currency=currency)

        # Cek apakah workbook asli dapat dimuat untuk menyalin sheet unselected (Requirement 4)
        wb = None
        if original_file_path and os.path.exists(original_file_path):
            ext_orig = os.path.splitext(original_file_path)[1].lower()
            if ext_orig in [".xlsx", ".xlsm"]:
                try:
                    wb = openpyxl.load_workbook(original_file_path, data_only=False)
                except Exception:
                    wb = None

        if wb is None:
            wb = openpyxl.Workbook()
            is_new_wb = True
        else:
            is_new_wb = False

        # Tulis/ganti sheet yang dibersihkan
        for idx, (s_name, s_df) in enumerate(sheet_dfs.items()):
            if is_new_wb:
                if idx == 0:
                    ws = wb.active
                    ws.title = str(s_name)
                else:
                    ws = wb.create_sheet(title=str(s_name))
            else:
                if str(s_name) in wb.sheetnames:
                    s_idx = wb.sheetnames.index(str(s_name))
                    del wb[str(s_name)]
                    ws = wb.create_sheet(title=str(s_name), index=s_idx)
                else:
                    ws = wb.create_sheet(title=str(s_name))

            ws.views.sheetView[0].showGridLines = True

            # Tulis Header
            for col_idx, col_name in enumerate(s_df.columns, start=1):
                ws.cell(row=1, column=col_idx, value=col_name)

            # Tulis Baris Data
            for row_idx, row_values in enumerate(s_df.itertuples(index=False), start=2):
                for col_idx, val in enumerate(row_values, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if pd.isna(val):
                        cell.value = None
                    elif hasattr(val, "to_pydatetime"):
                        cell.value = val.to_pydatetime()
                    else:
                        cell.value = val

            # Terapkan styling profesional
            styler.style_worksheet(ws, s_df, add_total_row=add_total_row)

        # Tambahkan Sheet Cleaning Summary jika diminta
        if include_summary_sheet:
            active_reports = sheet_reports if is_multi else (report or self.last_report)
            if active_reports:
                # Hapus sheet summary lama jika ada dari file asli
                for old_summary_title in ["Cleaning Summary", "Ringkasan Cleaning"]:
                    if old_summary_title in wb.sheetnames:
                        del wb[old_summary_title]

                reporter = CleaningReporter(theme=theme)
                if is_multi:
                    combined_data = {
                        s: {"report": sheet_reports.get(s, {}), "df": sheet_dfs.get(s)}
                        for s in sheet_dfs
                    }
                    reporter.add_summary_sheet(wb, combined_data)
                else:
                    reporter.add_summary_sheet(wb, active_reports, df)

        if orphan_df is not None:
            if "Baris Ter-drop" in wb.sheetnames:
                del wb["Baris Ter-drop"]
            ws_orphan = wb.create_sheet(title="Baris Ter-drop")
            ws_orphan.views.sheetView[0].showGridLines = True
            for col_idx, col_name in enumerate(orphan_df.columns, start=1):
                ws_orphan.cell(row=1, column=col_idx, value=col_name)
            for row_idx, row_values in enumerate(orphan_df.itertuples(index=False), start=2):
                for col_idx, value in enumerate(row_values, start=1):
                    cell = ws_orphan.cell(row=row_idx, column=col_idx)
                    cell.value = None if pd.isna(value) else value
            styler.style_worksheet(ws_orphan, orphan_df, add_total_row=False)

        if skipped_append_df is not None:
            if "Baris Dilewati" in wb.sheetnames:
                del wb["Baris Dilewati"]
            ws_skipped = wb.create_sheet(title="Baris Dilewati")
            ws_skipped.views.sheetView[0].showGridLines = True
            for col_idx, col_name in enumerate(skipped_append_df.columns, start=1):
                ws_skipped.cell(row=1, column=col_idx, value=col_name)
            for row_idx, row_values in enumerate(skipped_append_df.itertuples(index=False), start=2):
                for col_idx, value in enumerate(row_values, start=1):
                    cell = ws_skipped.cell(row=row_idx, column=col_idx)
                    cell.value = None if pd.isna(value) else value
            styler.style_worksheet(ws_skipped, skipped_append_df, add_total_row=False)

        wb.save(output_path)
        print(f"💾 File Excel profesional berhasil disimpan ke: {output_path}")
        return output_path

